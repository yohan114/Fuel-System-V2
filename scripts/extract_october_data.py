import openpyxl
import json
import sys

sys.stdout.reconfigure(encoding='utf-8')

path = r'C:\Users\HP\.gemini\antigravity\brain\9d7dc2ec-39ac-46c9-9538-f7fae859cb50\.user_uploaded\media_1791171656413.xlsx'
wb = openpyxl.load_workbook(path, data_only=True)
ws = wb['Oct']

items = []

for r in range(5, 185):
    if r in [100, 184]:  # Skip subtotal rows
        continue
    reg = str(ws.cell(r, 2).value or '').strip()
    code = str(ws.cell(r, 3).value or '').strip()
    vtype = str(ws.cell(r, 4).value or '').strip()
    total_tank = ws.cell(r, 37).value or 0
    
    days = {}
    for d in range(1, 32):
        col = d + 5
        v = ws.cell(r, col).value
        if v is not None and v != 0:
            days[d] = float(v)
            
    if days or total_tank > 0:
        items.append({
            'row': r,
            'reg': reg,
            'code': code,
            'type': vtype,
            'total_tank': float(total_tank),
            'days': days,
            'opening_meter': ws.cell(r, 40).value,
            'closing_meter': ws.cell(r, 41).value,
            'running_meter': ws.cell(r, 42).value,
        })

print(f"Extracted {len(items)} active items from October sheet.")
total_vol = sum(sum(it['days'].values()) for it in items)
print(f"Total volume: {total_vol} Litres")

with open('scripts/october_parsed.json', 'w', encoding='utf-8') as f:
    json.dump({
        'month': '2026-10',
        'total_volume': total_vol,
        'items': items,
        'receipts': [
            {'day': 1, 'litres': 2000.0},
            {'day': 2, 'litres': 200.0},
            {'day': 3, 'litres': 2100.0}
        ],
        'closing_balance': 603.0
    }, f, indent=2)

print("Saved to scripts/october_parsed.json")
